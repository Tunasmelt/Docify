import base64
import logging
import os
import re
from dataclasses import dataclass, field
from enum import Enum
from io import BytesIO

import docx
from docx.table import Table as DocxTable
import httpx
import pdfplumber
import pptx
from pptx.enum.shapes import PP_PLACEHOLDER
import pytesseract
from google import genai
from google.genai import types
from PIL import Image
from selectolax.parser import HTMLParser

logger = logging.getLogger(__name__)

# FEAT-017: gemini-2.5-flash, not the 3.6/3.5 models used for generation/
# verification — this is Flash proper (vision-native, free tier 1,500
# req/day), unrelated to the generate/verify model choice. Live-confirmed
# callable against the real API before use (2026-07-25), not assumed from
# .agent/api-docs/gemini.md's model table alone.
OCR_MODEL = "gemini-2.5-flash"

OCR_SYSTEM_PROMPT = (
    "Transcribe all readable text from this scanned document page, in reading "
    "order, as plain text. Do not describe the image or add commentary — output "
    "only the transcribed text. If the page has no readable text at all, output "
    "nothing."
)

# Audit finding (2026-07-26): with no explicit http_options.timeout, the
# genai SDK passes timeout=None straight through to httpx — confirmed
# directly against the installed SDK source (_api_client.py), not
# assumed — which httpx treats as "no timeout at all," not "use a
# default." A hung connection would block this call, and therefore the
# whole per-page OCR step (and the pages after it), indefinitely. Same
# reasoning as OcrSpaceClient's explicit 60s below.
OCR_TIMEOUT_MS = 60_000


class GeminiOcrClient:
    """Tier 1 of FEAT-017's OCR fallback chain — a page whose real parse
    yielded suspiciously little gets sent here first, as a rendered image,
    for a real vision-model transcription. Never raises: a failed call
    returns None so the chain (in Parser.parse()) moves on to tier 2,
    matching this project's established fail-safe discipline elsewhere
    (Verifier's fail-to-unsupported pattern)."""

    def __init__(self, client: genai.Client | None = None):
        # Deliberately NOT resolved here (audit finding, 2026-07-26): the
        # original eager `client or genai.Client(api_key=os.environ[...])`
        # read GEMINI_API_KEY at Parser()-construction time — meaning a
        # missing key crashed the whole Parser(), even for a document that
        # would never once trigger OCR. Resolved lazily instead, inside
        # transcribe_page()'s own try/except below, so a missing/bad key
        # becomes an ordinary per-call tier failure (logged, chain moves
        # to tier 2) — never a Parser()-construction-time crash.
        self._client = client

    def _get_client(self) -> genai.Client:
        if self._client is None:
            self._client = genai.Client(
                api_key=os.environ["GEMINI_API_KEY"],
                http_options=types.HttpOptions(timeout=OCR_TIMEOUT_MS),
            )
        return self._client

    def transcribe_page(self, image: Image.Image) -> str | None:
        try:
            client = self._get_client()
            buf = BytesIO()
            image.save(buf, format="PNG")
            response = client.models.generate_content(
                model=OCR_MODEL,
                contents=[
                    types.Part.from_text(text=OCR_SYSTEM_PROMPT),
                    types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png"),
                ],
            )
            text = response.text
            return text.strip() if text and text.strip() else None
        except Exception:
            logger.warning("parser: Gemini OCR tier call failed for a page", exc_info=True)
            return None


# Tier 2: OCR.space — a plain REST API, not an SDK (.agent/api-docs/ocrspace.md,
# verified live 2026-07-26). Deliberately a second, independent vendor: a
# Gemini-side outage or quota exhaustion (a real, hit-live constraint —
# see .agent/MEMORY.md's 2026-07-26 entry) has zero chance of also taking
# out this tier, since it's a different company's infrastructure entirely.
OCR_SPACE_URL = "https://api.ocr.space/parse/image"
OCR_SPACE_ENGINE = 2  # the newer/more-accurate of OCR.space's two engines — see ocrspace.md


class OcrSpaceClient:
    """Tier 2. Same never-raises contract as GeminiOcrClient — a failure
    here (network, auth, or the API's own IsErroredOnProcessing flag,
    which can come back on a 200 OK) returns None so the chain falls
    through to tier 3."""

    def __init__(self, api_key: str | None = None, http_client: httpx.Client | None = None):
        # api_key resolution deliberately deferred to transcribe_page()'s
        # own try/except (same reasoning as GeminiOcrClient, audit finding
        # 2026-07-26) — os.environ["OCR_SPACE_API_KEY"] here at
        # construction time would crash Parser() itself if unset, not just
        # this one tier. The httpx.Client itself is safe to build eagerly;
        # it doesn't need the key.
        self._api_key = api_key
        self._http = http_client or httpx.Client(timeout=OCR_TIMEOUT_MS / 1000.0)

    def _get_api_key(self) -> str:
        if self._api_key is None:
            self._api_key = os.environ["OCR_SPACE_API_KEY"]
        return self._api_key

    def transcribe_page(self, image: Image.Image) -> str | None:
        try:
            api_key = self._get_api_key()
            buf = BytesIO()
            image.save(buf, format="PNG")
            b64 = base64.b64encode(buf.getvalue()).decode()
            response = self._http.post(
                OCR_SPACE_URL,
                headers={"apikey": api_key},
                data={"base64Image": f"data:image/png;base64,{b64}", "OCREngine": OCR_SPACE_ENGINE},
            )
            response.raise_for_status()
            body = response.json()
            # A processing failure comes back as a normal 200 OK with this
            # flag set — raise_for_status() above does not catch it.
            if body.get("IsErroredOnProcessing"):
                logger.warning("parser: OCR.space reported a processing error: %s", body.get("ErrorMessage"))
                return None
            parsed_results = body.get("ParsedResults") or []
            if not parsed_results:
                return None
            text = parsed_results[0].get("ParsedText")
            return text.strip() if text and text.strip() else None
        except Exception:
            logger.warning("parser: OCR.space tier call failed for a page", exc_info=True)
            return None


class TesseractOcrClient:
    """Tier 3, the last resort: self-hosted, no network call, no vendor
    quota of any kind to exhaust — always available as long as the
    container has the tesseract-ocr system binary installed (Docker-only
    on Render; see ARCHITECTURE.md's deploy constraint). TESSERACT_CMD
    lets local dev point at a binary that isn't on PATH (e.g. a Windows
    install) without affecting the Docker/Linux production path, where
    `tesseract` is already on PATH after the apt-get install."""

    def __init__(self, tesseract_cmd: str | None = None):
        cmd = tesseract_cmd or os.environ.get("TESSERACT_CMD")
        if cmd:
            pytesseract.pytesseract.tesseract_cmd = cmd

    def transcribe_page(self, image: Image.Image) -> str | None:
        try:
            # timeout=0 (pytesseract's own default) means "no timeout at
            # all" — confirmed directly against the installed source
            # (pytesseract.py's timeout_manager: a falsy value skips
            # subprocess.communicate()'s own timeout entirely). A hung
            # tesseract process (a real, documented failure mode on
            # certain pathological images) would otherwise block this
            # call, and the rest of the parse, indefinitely.
            text = pytesseract.image_to_string(image, timeout=OCR_TIMEOUT_MS // 1000)
            return text.strip() if text and text.strip() else None
        except RuntimeError as exc:
            logger.warning("parser: Tesseract OCR tier timed out for a page: %s", exc)
            return None
        except Exception:
            logger.warning("parser: Tesseract OCR tier call failed for a page", exc_info=True)
            return None


def _default_ocr_tiers() -> list[tuple[str, object]]:
    return [
        ("gemini", GeminiOcrClient()),
        ("ocrspace", OcrSpaceClient()),
        ("tesseract", TesseractOcrClient()),
    ]


# Data contract lives in services/document_model.py (FEAT-027) — that
# module has no dependency on the heavy parsing libraries imported above,
# so chunker.py (and everything that imports chunker.py) can depend on
# ElementType/ParsedElement/ParsedDocument at runtime without paying for
# pdfplumber/docx/pptx/pytesseract/selectolax/google.genai. Re-exported
# here so `from services.parser import ElementType` etc. (existing call
# sites, tests) keeps working unchanged.
from services.document_model import (  # noqa: E402
    _TEXTUAL_ELEMENT_TYPES,
    BBox,
    ElementType,
    ParsedDocument,
    ParsedElement,
    ParseError,
)


# ══════════════════════════════════════════════════════════════════════════
# Shared heuristics (FEAT-027, .agent/reviews/2026-08-01-parser-research.md)
# ══════════════════════════════════════════════════════════════════════════
#
# PDF is the only format needing any of this — DOCX/PPTX/HTML all expose
# real, native structural metadata (paragraph styles, placeholder types,
# semantic tags) directly through their own libraries, confirmed live
# against every real fixture in the research pass. Nothing below is used
# for those three formats.

# Caption prefix — matches Docling-labeled captions' own real text shape in
# every fixture checked ("Table 1", "Table 2: example of...", "Figure 1: a
# real embedded PNG chart image."). Anchored to line start (not
# `re.search`) so it never matches a mid-sentence reference like "see
# Table 4 for details" — AND requires a colon or end-of-string immediately
# after the number, not just any following word. A real, live false
# positive found during implementation: "Table 1 below shows quarterly
# revenue figures..." (a body sentence that merely starts with "Table 1")
# matched a looser version of this regex in both the DOCX and HTML
# fixtures, which do carry that exact sentence alongside the real caption.
# Every genuine caption in every fixture checked is either "Table N:
# description" or a bare "Table N" — never "Table N <more prose>".
_CAPTION_PREFIX_RE = re.compile(r"^(Table|Figure)\s+\d+\s*(:|$)", re.IGNORECASE)

# Bullet/number list-item prefix. Catches numbered ("1. ", "2) "),
# lettered ("a. "), and parenthesized-footnote ("(1) ") lists directly
# from extracted text — the last of these a real gap found live against
# table_heavy.pdf: its own 6 Docling-labeled list items are all
# parenthesized footnote markers ("(1) Provisional total as of publication
# date."), a distinct shape from clean_digital.pdf's plain "1. "/"2. "
# numbering that an earlier version of this regex didn't cover.
# Deliberately does NOT catch every unordered list on its own — research
# found a real, confirmed gap: a PDF's unordered list can carry no bullet
# glyph at all in its text layer (only visual indentation), which this
# regex alone cannot see. The indent-based check below closes that gap.
_BULLET_PREFIX_RE = re.compile(r"^\s*([•\-\*•●▪]|\(\d+\)|\d+[.)]|[a-z][.)])\s+")

# A SHORT line indented further than the document's own body-text baseline
# by at least this many points is treated as a list item even with no
# bullet glyph — confirmed against clean_digital.pdf's real geometry: body
# text sits at x0≈43.3, its glyph-less unordered list at x0≈71.5 (indent
# +28.2). Deliberately set ABOVE the fixture's own indented-but-NOT-a-list
# content (a block-quote paragraph at x0≈61.0, indent +17.7) rather than
# just above zero — a real, live false positive found during
# implementation: indent alone also fires on a quoted paragraph and on a
# deeply-indented trailing credit line, neither of which is a list. The
# word-count guard below is what actually separates a real list item
# ("Item 1", 2 words) from indented prose (a full sentence) — indent alone
# is necessary but not sufficient.
_LIST_INDENT_THRESHOLD_PT = 20.0
_LIST_INDENT_MAX_WORDS = 5

# A whole-line-bold, short run is treated as a heading even at body-text
# size or smaller — confirmed against clean_digital.pdf: "Lists" (bold,
# same size as body), "Quote" and "Table" (bold, SMALLER than body — a
# real small-caps-style section label) are all headings Docling itself
# labeled, none of which the size-based rule alone catches. Guarded by
# word count so a genuinely bold PARAGRAPH (not a short label) is never
# misclassified — `_group_chars_into_lines`' `is_bold` is already
# whole-line (every character bold), so a bold RUN inside an otherwise
# plain sentence never trips this either.
_HEADING_BOLD_MAX_WORDS = 6

# An image covering more of the page than this is treated as "this page IS
# a scan," not a meaningful embedded figure worth extracting on its own —
# confirmed live against scanned.pdf: pdfplumber's page.images reports one
# image per page covering ~100% of the page area on all 3 pages (the raw
# scan itself), categorically different from a real embedded photo/logo/
# chart. Extracting the whole-page scan as a "figure" would be redundant
# with what the OCR fallback below already recovers as text for that same
# page, not genuine additional content.
_FIGURE_MAX_PAGE_COVERAGE = 0.85

# A detected table with fewer than this many rows is treated as a false
# positive, not a real table — confirmed live: pdfplumber's geometric table
# detector misreads a single-row bordered/indented block (e.g. a styled
# block-quote paragraph in clean_digital.pdf) as a 1-row "table." A real
# table in every fixture checked has a header row plus at least one data
# row.
_MIN_TABLE_ROWS = 2


def _rows_to_markdown(rows: list[list[str | None]]) -> str:
    if not rows:
        return ""
    header = rows[0]
    body = rows[1:]

    def cell(value: str | None) -> str:
        return (value or "").replace("\n", " ").replace("|", "\\|").strip()

    lines = ["| " + " | ".join(cell(c) for c in header) + " |"]
    lines.append("|" + "|".join(["---"] * len(header)) + "|")
    for row in body:
        lines.append("| " + " | ".join(cell(c) for c in row) + " |")
    return "\n".join(lines)


def _bbox_overlaps(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> bool:
    ax0, atop, ax1, abottom = a
    bx0, btop, bx1, bbottom = b
    return not (ax1 < bx0 or ax0 > bx1 or abottom < btop or atop > bbottom)


def _bbox_center(bbox: tuple[float, float, float, float]) -> tuple[float, float]:
    x0, top, x1, bottom = bbox
    return ((x0 + x1) / 2, (top + bottom) / 2)


def _bbox_distance(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ax, ay = _bbox_center(a)
    bx, by = _bbox_center(b)
    return ((ax - bx) ** 2 + (ay - by) ** 2) ** 0.5


# ══════════════════════════════════════════════════════════════════════════
# PDF extraction (pdfplumber)
# ══════════════════════════════════════════════════════════════════════════


def _document_body_font_size(pdf) -> float:
    """The single most common (char-count-weighted) font size across the
    whole document — the body-text baseline every heading heuristic below
    is judged against. Computed once per document, not per page: real
    fixtures checked (clean_digital.pdf, table_heavy.pdf) both have one
    dominant body size used throughout, with headings a clear step above
    it — confirmed live, not assumed."""
    sizes: dict[float, int] = {}
    for page in pdf.pages:
        for ch in page.chars:
            size = round(ch["size"], 1)
            sizes[size] = sizes.get(size, 0) + 1
    if not sizes:
        return 11.0  # no text at all on any page — fallback, never divides by zero below
    return max(sizes, key=lambda s: sizes[s])


def _group_chars_into_lines(page) -> list[dict]:
    """Groups a page's characters into visual lines by vertical position —
    pdfplumber's own primitive (page.chars) has no line concept built in.
    Each line carries its text, bbox, average font size, and whether every
    character in it is bold (by font name — pdfplumber doesn't expose a
    separate bold flag)."""
    lines: dict[float, list] = {}
    for ch in page.chars:
        key = round(ch["top"], 0)
        lines.setdefault(key, []).append(ch)

    result = []
    for top in sorted(lines.keys()):
        chars = sorted(lines[top], key=lambda c: c["x0"])
        text = "".join(c["text"] for c in chars).strip()
        if not text:
            continue
        x0 = chars[0]["x0"]
        x1 = chars[-1]["x1"]
        bottom = max(c["bottom"] for c in chars)
        avg_size = sum(c["size"] for c in chars) / len(chars)
        is_bold = all("bold" in c["fontname"].lower() for c in chars)
        result.append(
            {
                "text": text,
                "top": top,
                "bottom": bottom,
                "x0": x0,
                "x1": x1,
                "size": avg_size,
                "bold": is_bold,
            }
        )
    return result


def _classify_line(line: dict, body_size: float, body_x0: float) -> ElementType:
    if _CAPTION_PREFIX_RE.match(line["text"]):
        return ElementType.CAPTION

    word_count = len(line["text"].split())

    # Heading: clearly larger than body text (any weight), OR whole-line
    # bold and short (catches a same-/smaller-size bold section label —
    # see _HEADING_BOLD_MAX_WORDS's docstring for the real fixture finding
    # this covers).
    if line["size"] > body_size * 1.08:
        return ElementType.HEADING
    if line["bold"] and word_count <= _HEADING_BOLD_MAX_WORDS:
        return ElementType.HEADING

    if _BULLET_PREFIX_RE.match(line["text"]):
        return ElementType.LIST
    if line["x0"] > body_x0 + _LIST_INDENT_THRESHOLD_PT and word_count <= _LIST_INDENT_MAX_WORDS:
        return ElementType.LIST
    return ElementType.TEXT


def _document_body_x0(pdf) -> float:
    """Left-margin baseline for the indent-based list heuristic — the most
    common line-start x0 across the document, mirroring
    _document_body_font_size's same weighting approach."""
    positions: dict[float, int] = {}
    for page in pdf.pages:
        for line in _group_chars_into_lines(page):
            key = round(line["x0"], 0)
            positions[key] = positions.get(key, 0) + 1
    if not positions:
        return 0.0
    return max(positions, key=lambda x: positions[x])


def _extract_pdf_tables(page) -> list[dict]:
    """Detected tables on one page, each with its real bbox and markdown
    content — false positives (fewer than _MIN_TABLE_ROWS rows) filtered
    out. Confirmed live: this filter removes exactly the one shared
    false-positive both candidate libraries produced in the research pass
    (a styled block-quote misread as a 1-row table) without affecting the
    29/29 real detection rate on table_heavy.pdf."""
    results = []
    for table in page.find_tables():
        rows = table.extract()
        if len(rows) < _MIN_TABLE_ROWS:
            continue
        results.append({"bbox": table.bbox, "markdown": _rows_to_markdown(rows)})
    return results


def _extract_pdf_figures(page) -> list[dict]:
    """Real embedded images only — page-covering "images" (a scanned
    page's own raw content, not a meaningful embedded figure) are filtered
    by _FIGURE_MAX_PAGE_COVERAGE. See that constant's docstring for the
    live finding this is based on."""
    page_area = float(page.width) * float(page.height)
    results = []
    for img in page.images:
        bbox = (img["x0"], img["top"], img["x1"], img["bottom"])
        area = max(0.0, img["x1"] - img["x0"]) * max(0.0, img["bottom"] - img["top"])
        if page_area > 0 and area / page_area > _FIGURE_MAX_PAGE_COVERAGE:
            continue
        try:
            cropped = page.crop(bbox)
            pil_image = cropped.to_image(resolution=150).original.convert("RGB")
        except Exception:
            logger.warning("parser: dropped a figure on page %s — image render failed", page.page_number, exc_info=True)
            continue
        results.append({"bbox": bbox, "image": pil_image})
    return results


def _match_captions_to_targets(
    captions: list[dict], targets: list[dict]
) -> dict[int, int | None]:
    """Tier 1 caption<->table/figure association (FEAT-027) — same
    reading-order-then-bbox-distance greedy algorithm chunker.py's own
    Tier 2 already uses (`_resolve_tier2_captions`), reused here at Tier 1
    since this parser has no ML-provided explicit link to prefer over it.
    Returns {caption_index: target_index_or_None}. A target already
    claimed by an earlier (in reading-order) caption is no longer a
    candidate for a later one — same documented greedy limitation
    chunker.py's own version has (.agent/MEMORY.md §Anti-patterns),
    carried forward deliberately for consistency, not independently
    re-decided here."""
    resolution: dict[int, int | None] = {}
    claimed: set[int] = set()
    for c_idx, caption in enumerate(captions):
        candidates = [t_idx for t_idx in range(len(targets)) if t_idx not in claimed]
        if not candidates:
            resolution[c_idx] = None
            continue
        best = min(candidates, key=lambda t_idx: _bbox_distance(caption["bbox"], targets[t_idx]["bbox"]))
        resolution[c_idx] = best
        claimed.add(best)
    return resolution


def _parse_pdf(file_bytes: bytes) -> tuple[list[ParsedElement], int]:
    last_page_number: int | None = None
    try:
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            if not pdf.pages:
                raise ParseError("PDF has no pages")
            body_size = _document_body_font_size(pdf)
            body_x0 = _document_body_x0(pdf)

            elements: list[ParsedElement] = []
            dropped = 0
            element_counter = 0

            for page in pdf.pages:
                page_number = page.page_number
                last_page_number = page_number

                tables = _extract_pdf_tables(page)
                figures = _extract_pdf_figures(page)
                target_bboxes = [t["bbox"] for t in tables] + [f["bbox"] for f in figures]

                lines = _group_chars_into_lines(page)
                non_table_lines = [
                    line
                    for line in lines
                    if not any(
                        _bbox_overlaps((line["x0"], line["top"], line["x1"], line["bottom"]), tb)
                        for tb in target_bboxes
                    )
                ]

                # Classify every remaining line; merge consecutive TEXT
                # lines into one paragraph-level element (small vertical
                # gap = still the same paragraph), matching the
                # granularity real documents' own paragraph structure
                # actually has — HEADING/LIST lines never merge, each
                # stays its own element. CAPTION lines DO merge with
                # immediately-following close-gap lines that don't
                # themselves start something new (audit finding,
                # 2026-08-02: a caption's own text was silently truncated
                # to its first visual line — e.g. "Table 10: ... (multiple"
                # dropped "layout problems)" — because only the
                # prefix-matching first line was ever captured; a caption
                # wrapping onto a second line is common and was never
                # continued).
                page_captions: list[dict] = []
                page_elements: list[dict] = []
                text_buffer: list[dict] = []
                caption_buffer: list[dict] = []

                def flush_text_buffer():
                    nonlocal element_counter
                    if not text_buffer:
                        return
                    combined_text = "\n".join(l["text"] for l in text_buffer)
                    bbox = (
                        min(l["x0"] for l in text_buffer),
                        text_buffer[0]["top"],
                        max(l["x1"] for l in text_buffer),
                        text_buffer[-1]["bottom"],
                    )
                    page_elements.append({"type": ElementType.TEXT, "bbox": bbox, "content": combined_text})
                    text_buffer.clear()

                def flush_caption_buffer():
                    if not caption_buffer:
                        return
                    combined_text = " ".join(l["text"] for l in caption_buffer)
                    bbox = (
                        min(l["x0"] for l in caption_buffer),
                        caption_buffer[0]["top"],
                        max(l["x1"] for l in caption_buffer),
                        caption_buffer[-1]["bottom"],
                    )
                    page_captions.append({"bbox": bbox, "text": combined_text})
                    caption_buffer.clear()

                prev_bottom = None
                for line in non_table_lines:
                    kind = _classify_line(line, body_size, body_x0)
                    line_height = line["bottom"] - line["top"] or 12.0
                    gap = (line["top"] - prev_bottom) if prev_bottom is not None else 0.0

                    # A close-gap line right after an in-progress caption is
                    # that caption's continuation — a wrapped caption line
                    # has no "Table N:" prefix of its own, so _classify_line
                    # reads it as either TEXT or (since it usually shares the
                    # caption's own bold styling — confirmed live: "layout
                    # problems)" continuing "Table 10: ... (multiple" is
                    # short + bold, matching the whole-line-bold heading
                    # heuristic) HEADING. Absorb both into the caption
                    # instead of starting a new element; LIST is excluded —
                    # a bullet-prefixed line is a genuine new list item, not
                    # a caption wrapping onto another line.
                    if (
                        caption_buffer
                        and kind in (ElementType.TEXT, ElementType.HEADING)
                        and gap <= line_height * 1.5
                    ):
                        caption_buffer.append(line)
                        prev_bottom = line["bottom"]
                        continue

                    if kind == ElementType.TEXT:
                        if text_buffer and gap > line_height * 1.5:
                            flush_text_buffer()
                        flush_caption_buffer()
                        text_buffer.append(line)
                    else:
                        flush_text_buffer()
                        if kind == ElementType.CAPTION:
                            flush_caption_buffer()
                            caption_buffer.append(line)
                        else:
                            flush_caption_buffer()
                            bbox = (line["x0"], line["top"], line["x1"], line["bottom"])
                            page_elements.append({"type": kind, "bbox": bbox, "content": line["text"]})
                    prev_bottom = line["bottom"]
                flush_text_buffer()
                flush_caption_buffer()

                # Tier 1: match this page's captions to this page's
                # tables/figures by reading-order + bbox proximity.
                targets = [{"bbox": t["bbox"], "kind": "table"} for t in tables] + [
                    {"bbox": f["bbox"], "kind": "figure"} for f in figures
                ]
                resolution = _match_captions_to_targets(page_captions, targets)

                # Build final ParsedElements for this page in real reading
                # order (sorted by vertical position) rather than grouped by
                # type — tables/figures/captions/text are interleaved the
                # way they actually appear on the page, matching Docling's
                # reading-order output instead of dumping every table before
                # any surrounding text.
                caption_ids_by_target_index: dict[int, list[str]] = {}
                caption_ids_by_c_idx: dict[int, str] = {}
                for c_idx, caption in enumerate(page_captions):
                    element_counter += 1
                    cap_id = f"pdf-p{page_number}-cap{element_counter}"
                    caption_ids_by_c_idx[c_idx] = cap_id
                    target_idx = resolution.get(c_idx)
                    if target_idx is not None:
                        caption_ids_by_target_index.setdefault(target_idx, []).append(cap_id)

                positioned: list[tuple[float, ParsedElement]] = []
                for t_idx, table in enumerate(tables):
                    element_counter += 1
                    positioned.append(
                        (
                            table["bbox"][1],
                            ParsedElement(
                                element_type=ElementType.TABLE,
                                page_number=page_number,
                                bbox=BBox(*table["bbox"]),
                                content=table["markdown"],
                                element_id=f"pdf-p{page_number}-table{element_counter}",
                                associated_caption_ids=caption_ids_by_target_index.get(t_idx, []),
                            ),
                        )
                    )
                for f_idx, figure in enumerate(figures):
                    element_counter += 1
                    positioned.append(
                        (
                            figure["bbox"][1],
                            ParsedElement(
                                element_type=ElementType.FIGURE,
                                page_number=page_number,
                                bbox=BBox(*figure["bbox"]),
                                content=figure["image"],
                                element_id=f"pdf-p{page_number}-fig{element_counter}",
                                associated_caption_ids=caption_ids_by_target_index.get(len(tables) + f_idx, []),
                            ),
                        )
                    )
                for c_idx, caption in enumerate(page_captions):
                    target_idx = resolution.get(c_idx)
                    method = "explicit" if target_idx is not None else "none"
                    positioned.append(
                        (
                            caption["bbox"][1],
                            ParsedElement(
                                element_type=ElementType.CAPTION,
                                page_number=page_number,
                                bbox=BBox(*caption["bbox"]),
                                content=caption["text"],
                                element_id=caption_ids_by_c_idx[c_idx],
                                association_method=method,
                            ),
                        )
                    )
                for el in page_elements:
                    element_counter += 1
                    positioned.append(
                        (
                            el["bbox"][1],
                            ParsedElement(
                                element_type=el["type"],
                                page_number=page_number,
                                bbox=BBox(*el["bbox"]),
                                content=el["content"],
                                element_id=f"pdf-p{page_number}-el{element_counter}",
                            ),
                        )
                    )

                positioned.sort(key=lambda item: item[0])
                elements.extend(el for _, el in positioned)

            return elements, dropped
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Failed to parse PDF: {exc}", page_number=last_page_number) from exc


def _render_pdf_page_image(file_bytes: bytes, page_number: int) -> Image.Image | None:
    """Renders one page as a full-resolution PIL Image for the OCR
    fallback below — done lazily, only for pages that actually need it
    (a fresh pdfplumber.open() per call is deliberate: cheap relative to
    an OCR API call, and avoids holding the whole PDF's page objects open
    for the entire parse just in case OCR is needed later)."""
    try:
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            if page_number < 1 or page_number > len(pdf.pages):
                return None
            page = pdf.pages[page_number - 1]
            return page.to_image(resolution=150).original.convert("RGB")
    except Exception:
        logger.warning("parser: failed to render page %s for OCR fallback", page_number, exc_info=True)
        return None


# ══════════════════════════════════════════════════════════════════════════
# DOCX extraction (python-docx)
# ══════════════════════════════════════════════════════════════════════════
#
# Real native structure, no heuristics needed: paragraph.style.name
# ("Heading 1"/"Heading 2"/... -> HEADING; "List Bullet"/"List Number" ->
# LIST; "Caption" -> CAPTION, confirmed live to round-trip correctly,
# .agent/reviews/2026-08-01-parser-research.md — table.docx's own fixture
# doesn't happen to use Word's Caption style, so this path is exercised by
# a synthetic construction in that research, and by whichever future real
# document does use it; a caption-shaped paragraph using the text-prefix
# convention below is caught regardless of style). Sentinel page_number=1
# and zero bbox for every element, matching the prior Docling-based
# contract exactly — DOCX has no real page/coordinate concept to report.

_DOCX_HEADING_STYLE_RE = re.compile(r"^Heading\s+\d+$", re.IGNORECASE)
_DOCX_LIST_STYLE_RE = re.compile(r"^List\b", re.IGNORECASE)
_SENTINEL_BBOX = BBox(x0=0.0, y0=0.0, x1=0.0, y1=0.0)


def _docx_image_bytes_by_rid(document, rid: str) -> bytes | None:
    try:
        return document.part.related_parts[rid].blob
    except Exception:
        return None


def _parse_docx(file_bytes: bytes) -> tuple[list[ParsedElement], int]:
    try:
        document = docx.Document(BytesIO(file_bytes))
    except Exception as exc:
        raise ParseError(f"Failed to open DOCX: {exc}") from exc

    elements: list[ParsedElement] = []
    dropped = 0
    counter = 0

    def add_figure(image_bytes: bytes | None) -> None:
        nonlocal counter, dropped
        if image_bytes is None:
            dropped += 1
            logger.warning("parser: dropped a DOCX figure — could not resolve its image bytes")
            return
        try:
            pil_image = Image.open(BytesIO(image_bytes)).convert("RGB")
            pil_image.load()  # force decode now, while the source bytes are still in scope
        except Exception:
            dropped += 1
            logger.warning("parser: dropped a DOCX figure — image bytes failed to decode", exc_info=True)
            return
        counter += 1
        elements.append(
            ParsedElement(
                element_type=ElementType.FIGURE,
                page_number=1,
                bbox=_SENTINEL_BBOX,
                content=pil_image,
                element_id=f"docx-fig{counter}",
            )
        )

    try:
        # Body order, not "all paragraphs, then all tables, then all images":
        # the chunker groups adjacent elements and matches captions by
        # proximity, so a table or figure must stay where it sits in the text.
        for block in document.iter_inner_content():
            if isinstance(block, DocxTable):
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                if len(rows) < _MIN_TABLE_ROWS:
                    continue
                counter += 1
                elements.append(
                    ParsedElement(
                        element_type=ElementType.TABLE,
                        page_number=1,
                        bbox=_SENTINEL_BBOX,
                        content=_rows_to_markdown(rows),
                        element_id=f"docx-table{counter}",
                    )
                )
                continue

            paragraph = block
            text = paragraph.text.strip()
            if text:
                style_name = paragraph.style.name if paragraph.style else ""
                if style_name.lower() == "caption" or _CAPTION_PREFIX_RE.match(text):
                    element_type = ElementType.CAPTION
                elif _DOCX_HEADING_STYLE_RE.match(style_name):
                    element_type = ElementType.HEADING
                elif _DOCX_LIST_STYLE_RE.match(style_name) or _BULLET_PREFIX_RE.match(text):
                    element_type = ElementType.LIST
                else:
                    element_type = ElementType.TEXT
                counter += 1
                elements.append(
                    ParsedElement(
                        element_type=element_type,
                        page_number=1,
                        bbox=_SENTINEL_BBOX,
                        content=text,
                        element_id=f"docx-p{counter}",
                    )
                )

            # Inline images live inside paragraph runs — same scope as the
            # old document.inline_shapes (floating/anchored images excluded).
            for rid in paragraph._element.xpath(".//wp:inline//a:blip/@r:embed"):
                add_figure(_docx_image_bytes_by_rid(document, rid))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Failed while processing DOCX elements: {exc}") from exc

    return elements, dropped


# ══════════════════════════════════════════════════════════════════════════
# PPTX extraction (python-pptx)
# ══════════════════════════════════════════════════════════════════════════
#
# Real native structure: placeholder type (CENTER_TITLE/TITLE/SUBTITLE ->
# HEADING) is a direct, zero-heuristic signal — cleaner than DOCX's style
# names or PDF's font geometry. page_number is the REAL 1-indexed slide
# index (confirmed live, FEAT-020) — PPTX is not a sentinel format, unlike
# DOCX/HTML. bbox uses the shape's real EMU position/size, converted to
# points (1 pt = 12700 EMU) to stay unit-consistent with the PDF path.

_EMU_PER_POINT = 12700


def _pptx_shape_bbox(shape) -> BBox:
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        if None in (left, top, width, height):
            return _SENTINEL_BBOX
        x0 = left / _EMU_PER_POINT
        y0 = top / _EMU_PER_POINT
        return BBox(x0=x0, y0=y0, x1=x0 + width / _EMU_PER_POINT, y1=y0 + height / _EMU_PER_POINT)
    except Exception:
        return _SENTINEL_BBOX


def _parse_pptx(file_bytes: bytes) -> tuple[list[ParsedElement], int]:
    try:
        presentation = pptx.Presentation(BytesIO(file_bytes))
    except Exception as exc:
        raise ParseError(f"Failed to open PPTX: {exc}") from exc

    elements: list[ParsedElement] = []
    dropped = 0
    counter = 0

    try:
        for slide_index, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                if shape.shape_type is not None and str(shape.shape_type) == "PICTURE (13)":
                    try:
                        image_bytes = shape.image.blob
                        pil_image = Image.open(BytesIO(image_bytes)).convert("RGB")
                        pil_image.load()
                    except Exception:
                        dropped += 1
                        logger.warning(
                            "parser: dropped a PPTX figure on slide %s — image extraction failed",
                            slide_index,
                            exc_info=True,
                        )
                        continue
                    counter += 1
                    elements.append(
                        ParsedElement(
                            element_type=ElementType.FIGURE,
                            page_number=slide_index,
                            bbox=_pptx_shape_bbox(shape),
                            content=pil_image,
                            element_id=f"pptx-s{slide_index}-fig{counter}",
                        )
                    )
                    continue

                if shape.has_table:
                    tbl = shape.table
                    rows = [[cell.text for cell in row.cells] for row in tbl.rows]
                    if len(rows) < _MIN_TABLE_ROWS:
                        continue
                    counter += 1
                    elements.append(
                        ParsedElement(
                            element_type=ElementType.TABLE,
                            page_number=slide_index,
                            bbox=_pptx_shape_bbox(shape),
                            content=_rows_to_markdown(rows),
                            element_id=f"pptx-s{slide_index}-table{counter}",
                        )
                    )
                    continue

                if not shape.has_text_frame:
                    continue

                # Exact enum-member comparison, not substring containment —
                # 2026-08-02 audit finding: `"TITLE" in str(placeholder_type)`
                # also matched SUBTITLE (python-pptx's str() for that member
                # is "SUBTITLE (4)", which contains "TITLE" as a substring),
                # misclassifying real slide subtitles — genuinely
                # descriptive/body text — as HEADING elements. CENTER_TITLE
                # is a real distinct title-role placeholder (used on title
                # slides) and stays included; SUBTITLE does not.
                is_title = shape.is_placeholder and shape.placeholder_format.type in (
                    PP_PLACEHOLDER.TITLE,
                    PP_PLACEHOLDER.CENTER_TITLE,
                )
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    if is_title:
                        element_type = ElementType.HEADING
                    elif paragraph.level and paragraph.level > 0:
                        element_type = ElementType.LIST
                    elif _BULLET_PREFIX_RE.match(text):
                        element_type = ElementType.LIST
                    elif _CAPTION_PREFIX_RE.match(text):
                        element_type = ElementType.CAPTION
                    else:
                        element_type = ElementType.TEXT
                    counter += 1
                    elements.append(
                        ParsedElement(
                            element_type=element_type,
                            page_number=slide_index,
                            bbox=_pptx_shape_bbox(shape),
                            content=text,
                            element_id=f"pptx-s{slide_index}-el{counter}",
                        )
                    )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Failed while processing PPTX elements: {exc}") from exc

    return elements, dropped


# ══════════════════════════════════════════════════════════════════════════
# HTML extraction (selectolax)
# ══════════════════════════════════════════════════════════════════════════
#
# Real native structure: semantic tags are ground truth, no heuristics
# needed. Sentinel page_number=1 and zero bbox for every element, same
# reasoning as DOCX — HTML has no page/coordinate concept either.

_HTML_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
_HTML_BLOCK_TAGS = _HTML_HEADING_TAGS | {"p", "li", "table", "figcaption", "caption"}
_HTML_SKIP_TAGS = {"script", "style", "noscript", "template", "head"}


def _iter_html_blocks(root):
    """Yields block elements in document order. A matched block is not
    descended into, so nested matches (a <p> inside an <li>, a <caption>
    inside a <table>) are never emitted twice — the outer block's own text
    already includes them. (`tree.css("h1, …, p")` returns nodes grouped by
    selector, not in document order, which detached every heading from its
    section.)"""
    stack = [root]
    while stack:
        node = stack.pop()
        tag = node.tag
        if tag in _HTML_BLOCK_TAGS and node is not root:
            yield node
            continue
        if tag in _HTML_SKIP_TAGS:
            continue
        children = []
        child = node.child
        while child is not None:
            if child.tag != "-text":
                children.append(child)
            child = child.next
        stack.extend(reversed(children))


def _parse_html(file_bytes: bytes) -> tuple[list[ParsedElement], int]:
    try:
        tree = HTMLParser(file_bytes)
    except Exception as exc:
        raise ParseError(f"Failed to open HTML: {exc}") from exc

    elements: list[ParsedElement] = []
    dropped = 0
    counter = 0

    def emit_text(node, tag: str) -> None:
        nonlocal counter
        text = node.text(deep=True, separator=" ").strip()
        if not text:
            return
        if tag in _HTML_HEADING_TAGS:
            element_type = ElementType.HEADING
        elif tag == "li":
            element_type = ElementType.LIST
        elif tag in ("figcaption", "caption") or _CAPTION_PREFIX_RE.match(text):
            element_type = ElementType.CAPTION
        else:
            element_type = ElementType.TEXT
        counter += 1
        elements.append(
            ParsedElement(
                element_type=element_type,
                page_number=1,
                bbox=_SENTINEL_BBOX,
                content=text,
                element_id=f"html-el{counter}",
            )
        )

    try:
        for node in _iter_html_blocks(tree.body or tree.root):
            tag = node.tag
            if tag == "table":
                # A <caption> belongs to its table but precedes it in the DOM
                # — emit it first so it stays adjacent, as the chunker expects.
                for caption in node.css("caption"):
                    emit_text(caption, "caption")
                rows = []
                for tr in node.css("tr"):
                    cells = [c.text(deep=True, separator=" ").strip() for c in tr.css("td, th")]
                    if cells:
                        rows.append(cells)
                if len(rows) < _MIN_TABLE_ROWS:
                    continue
                counter += 1
                elements.append(
                    ParsedElement(
                        element_type=ElementType.TABLE,
                        page_number=1,
                        bbox=_SENTINEL_BBOX,
                        content=_rows_to_markdown(rows),
                        element_id=f"html-table{counter}",
                    )
                )
                continue

            emit_text(node, tag)

        # HTML figure extraction (img tags): confirmed real gap during the
        # original FEAT-020 investigation (Docling itself never resolved
        # <img> content to a usable image either) — this rewrite doesn't
        # fetch/decode <img> src content (would need a network fetch for a
        # remote src, or a relative-path filesystem read for a local one,
        # neither of which this parser has enough context to do safely).
        # Text/table extraction for HTML is unaffected; this is a known,
        # unchanged limitation, not a regression.
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Failed while processing HTML elements: {exc}") from exc

    return elements, dropped


# ══════════════════════════════════════════════════════════════════════════
# Parser — format dispatch + OCR fallback
# ══════════════════════════════════════════════════════════════════════════

# A page whose text layer is this short AND that is mostly one image is a
# scan with a stray text layer (a title, a stamp) — OCR it. scanned.pdf's
# page 1 is the real case: 13 characters over a full-page scan.
_LOW_YIELD_MAX_CHARS = 200
_SCANNED_PAGE_MIN_IMAGE_COVERAGE = 0.5
_OCR_REPLACEABLE_TYPES = {ElementType.TEXT, ElementType.HEADING, ElementType.LIST, ElementType.CAPTION}


def _largest_image_coverage(page) -> float:
    page_area = float(page.width) * float(page.height)
    if page_area <= 0:
        return 0.0
    largest = 0.0
    for img in page.images:
        area = max(0.0, img["x1"] - img["x0"]) * max(0.0, img["bottom"] - img["top"])
        largest = max(largest, area / page_area)
    return largest


_EXTENSION_TO_FORMAT = {
    "pdf": "pdf",
    "docx": "docx",
    "pptx": "pptx",
    "html": "html",
    "htm": "html",
}


class Parser:
    def __init__(self, ocr_tiers: list[tuple[str, object]] | None = None):
        # Real 3-tier chain by default — OCR fallback is production
        # behavior, not an opt-in extra a caller has to remember to wire
        # up. routes/ingest.py constructs Parser() with no arguments and
        # gets all three tiers automatically. Tests that don't want real
        # calls inject their own (name, fake) tier list.
        self._ocr_tiers = ocr_tiers if ocr_tiers is not None else _default_ocr_tiers()

    def parse(self, file_bytes: bytes, filename: str = "document.pdf") -> ParsedDocument:
        """filename drives format detection (extension-based — this parser
        does not content-sniff). A caller ingesting DOCX/PPTX/HTML MUST
        pass the real filename (or at least the real extension)."""
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        fmt = _EXTENSION_TO_FORMAT.get(ext)
        if fmt is None:
            raise ParseError(f"Unsupported file format: {filename!r}")

        if fmt == "pdf":
            elements, dropped = _parse_pdf(file_bytes)
        elif fmt == "docx":
            elements, dropped = _parse_docx(file_bytes)
        elif fmt == "pptx":
            elements, dropped = _parse_pptx(file_bytes)
        else:
            elements, dropped = _parse_html(file_bytes)

        if fmt == "pdf":
            elements = self._run_ocr_fallback(elements, file_bytes)

        return ParsedDocument(elements=elements, dropped_elements=dropped)

    def _run_ocr_fallback(self, elements: list[ParsedElement], file_bytes: bytes) -> list[ParsedElement]:
        """FEAT-017: OCR fallback for low-yield PDF pages. Re-derived
        against this parser's own real silent-failure behavior (research
        item 8) — pdfplumber, unlike Docling, never raises and never
        returns a page with zero pages total; a page pdfplumber found no
        text on simply produces zero TEXT/HEADING/TABLE/LIST elements for
        that page_number, the same trigger shape the original Docling-era
        heuristic already used. PPTX/DOCX/HTML never reach this method at
        all (called only from the fmt == "pdf" branch above) — mirrors the
        prior contract exactly: OCR fallback is a PDF-specific concept,
        confirmed structurally correct for the other three formats in the
        original FEAT-020 investigation and unchanged here.
        """
        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            page_info = {
                page.page_number: (float(page.width), float(page.height), _largest_image_coverage(page))
                for page in pdf.pages
            }

        textual_chars_by_page: dict[int, int] = {}
        for e in elements:
            if e.element_type in _TEXTUAL_ELEMENT_TYPES and isinstance(e.content, str):
                textual_chars_by_page[e.page_number] = textual_chars_by_page.get(e.page_number, 0) + len(e.content)

        for page_number, (width, height, image_coverage) in page_info.items():
            textual_chars = textual_chars_by_page.get(page_number)
            if textual_chars is not None and not (
                textual_chars < _LOW_YIELD_MAX_CHARS and image_coverage >= _SCANNED_PAGE_MIN_IMAGE_COVERAGE
            ):
                continue  # a real text layer — no OCR needed

            page_image = _render_pdf_page_image(file_bytes, page_number)
            if page_image is None:
                logger.warning(
                    "parser: page %s is low-yield and could not be rendered for OCR fallback",
                    page_number,
                )
                continue

            recovered_text: str | None = None
            recovered_tier = "none"
            for tier_name, tier_client in self._ocr_tiers:
                try:
                    text = tier_client.transcribe_page(page_image)
                except Exception:
                    logger.warning(
                        "parser: OCR tier=%s raised for page %s — trying next tier",
                        tier_name,
                        page_number,
                        exc_info=True,
                    )
                    continue
                if text and text.strip():
                    recovered_text = text.strip()
                    recovered_tier = tier_name
                    break
                logger.warning(
                    "parser: OCR tier=%s found no recoverable text on page %s — trying next tier",
                    tier_name,
                    page_number,
                )

            if not recovered_text:
                logger.warning(
                    "parser: OCR fallback exhausted all tiers for page %s — page remains low-yield", page_number
                )
                continue

            if textual_chars is not None:
                # A partly scanned page: OCR transcribes the whole page, including
                # the few characters of text layer, so the OCR text replaces the
                # page's text elements rather than duplicating them. Tables and
                # figures detected on the page are kept.
                elements = [
                    e
                    for e in elements
                    if not (e.page_number == page_number and e.element_type in _OCR_REPLACEABLE_TYPES)
                ]

            elements.append(
                ParsedElement(
                    element_type=ElementType.TEXT,
                    page_number=page_number,
                    bbox=BBox(x0=0.0, y0=0.0, x1=width, y1=height),
                    content=recovered_text,
                    element_id=f"ocr-page-{page_number}",
                )
            )
            logger.info("parser: OCR fallback recovered text on page %s via tier=%s", page_number, recovered_tier)

        # Stable sort: OCR elements land at the end of their own page, not at
        # the end of the document.
        elements.sort(key=lambda e: e.page_number)
        return elements
